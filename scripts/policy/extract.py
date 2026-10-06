"""חילוץ טווחי מדיניות לכל מסלול ממסמכים שנמשכו (policy/docs_index.json).

הרצה: python -m scripts.policy.extract [--all]   (ברירת מחדל: רק מסמכים שלא נותחו)
פלט: policy/tracks_policy.csv + .json - שורה לכל (חברה, מסמך, מסלול).

היוריסטיקה על עברית חופשית: מזהה כותרת מסלול ("מסלול ..."), ובחלון הטקסט עד
הכותרת הבאה מחלץ טווחי חשיפה (מינ'-מקס') לפי אפיק. כל שורה נושאת snippet מקור
וציון ביטחון - חובה לעבור ידנית על מסמכים ראשונים כי פורמט משתנה בין חברות.
"""
from urllib.parse import unquote
import argparse, csv, json, os, re, sys
from pathlib import Path
from datetime import datetime

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / (__import__("os").environ.get("POLICY_OUT") or "policy")  # ריצה לחברה: policy/companies/<LegalId>

PARSER_VERSION = 23  # (v23: עמודת "ליום" המאוחרת) (v22: "48.21%" לא 4821) (v21: בלי 'נכון לתאריך') (v20: תאריך כתוב בגוף המסמך, יומן שינויים לפי שמות עמודות) (v10: Ayalon, dated current column) # הגדלה = פרסור מחדש של כל המסמכים בריצה הבאה (שינוי בפרסרים)

NUM = r"(\d{1,3}(?:\.\d+)?)"
PCT = NUM + r"\s*%?"
# אפיק -> ביטוי זיהוי בטקסט
ASSETS = {
    "equity": r"(?:חשיפה\s+ל)?מניות",
    "bonds": r"(?:איגרות|אגרות)\s+חוב|אג[\"״']?ח",
    "foreign": r"חו[\"״]ל|חוץ\s+לארץ|נכסים\s+זרים",
    "fx": r"מט[\"״]ח|מטבע\s+חוץ",
    "illiquid": r"לא\s+סחיר|בלתי\s+סחיר",
}
TRACK_HDR = re.compile(r"^\s*(?:\d+[.)]\s*)?((?:מסלול|קרן\s+השתלמות|קופת\s+גמל|קרן\s+פנסיה)[^\n:]{3,80})$")


def parse_range(s: str):
    """'10%-25%' / '10 עד 25' / 'עד 30%' / 'לפחות 50%' -> (min, max) (None = לא צוין)."""
    m = re.search(PCT + r"\s*(?:-|–|—|עד|ל-?)\s*" + PCT, s)
    if m:
        lo, hi = float(m.group(1)), float(m.group(2))
        return (min(lo, hi), max(lo, hi))
    m = re.search(r"(?:עד|מקסימום|לכל\s+היותר|לא\s+יעלה\s+על)\s*" + PCT, s)
    if m:
        return (None, float(m.group(1)))
    m = re.search(r"(?:לפחות|מינימום|לא\s+פחות\s+מ)\s*" + PCT, s)
    if m:
        return (float(m.group(1)), None)
    m = re.search(PCT, s)
    return (float(m.group(1)),) * 2 if m and "%" in s else (None, None)


def norm_name(s: str) -> str:
    # "מחוג מסלול כללי - עדכון מיום 11.8.26": אותו מסלול כמו בגרסת ינואר - בלי סיומת העדכון
    s = re.sub(r"\s*[-–—,]?\s*\(?(?:עדכון|עודכן|מעודכן)\s+(?:מיום|ביום|ליום|מתאריך)\s*[\d./\-\s]+\)?\s*$", "", s)
    s = re.sub(r"[\"'״׳\-–—()\[\]:.,]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def extract_tracks_from_text(text: str) -> list[dict]:
    lines = [l.rstrip() for l in text.splitlines()]
    hdrs = [(i, norm_name(m.group(1))) for i, l in enumerate(lines) if (m := TRACK_HDR.match(l))]
    out = []
    for n, (i, name) in enumerate(hdrs):
        end = hdrs[n + 1][0] if n + 1 < len(hdrs) else len(lines)
        block = lines[i + 1:end][:60]
        rec = {"track_name": name, "snippet": " ".join(block)[:400], "confidence": "text"}
        for k, pat in ASSETS.items():
            for l in block:
                if re.search(pat, l) and re.search(r"\d\s*%|%\s*\d|\d+\s*(?:-|–|עד)\s*\d+", l):
                    lo, hi = parse_range(re.split(pat, l, maxsplit=1)[-1])
                    if lo is not None or hi is not None:
                        rec[f"{k}_min"], rec[f"{k}_max"] = lo, hi
                        break
        for l in block:
            if re.search(r"מדד\s+ה?ייחוס|בנצ[׳']?מרק|benchmark", l, re.I):
                rec["benchmark"] = l.strip()[:200]; break
        if any(f"{k}_{b}" in rec for k in ASSETS for b in ("min", "max")):
            out.append(rec)
    return out


def extract_tracks_from_table(rows: list[list[str]]) -> list[dict]:
    """טבלה: שורת כותרת עם שמות אפיקים, שורה לכל מסלול (או להפך - לא נתמך עדיין)."""
    out, hdr = [], None
    for r in rows:
        cells = [(c or "").replace("\n", " ").strip() for c in r]
        if hdr is None:
            cols = {}
            for j, c in enumerate(cells):
                for k, pat in ASSETS.items():
                    if re.search(pat, c):
                        cols.setdefault(k, j)
            if len(cols) >= 2:
                hdr = cols
            continue
        if not cells or not cells[0] and len(cells) < 2:
            continue
        name = next((c for c in cells if c and not re.search(r"\d\s*%", c)), "")
        if not name:
            continue
        rec = {"track_name": norm_name(name), "snippet": " | ".join(cells)[:400], "confidence": "table"}
        for k, j in hdr.items():
            if j < len(cells) and cells[j]:
                lo, hi = parse_range(cells[j] if "%" in cells[j] else cells[j] + "%")
                if lo is not None or hi is not None:
                    rec[f"{k}_min"], rec[f"{k}_max"] = lo, hi
        if any(f"{k}_{b}" in rec for k in ASSETS for b in ("min", "max")):
            out.append(rec)
    return out


_HEB = re.compile(r"[\u0590-\u05FF]")


def _is_visual_rtl(text: str) -> bool:
    """PDF בעברית חזותית: אותיות סופיות (ךםןףץ) בתחילת מילים במקום בסופן."""
    words = re.findall(r"[\u0590-\u05FF]{2,}", text)
    if len(words) < 10:
        return False
    start = sum(w[0] in "ךםןףץ" for w in words)
    end = sum(w[-1] in "ךםןףץ" for w in words)
    return start > end


def _fix_rtl(s: str) -> str:
    """היפוך סדר המילים בשורה והיפוך תווים רק במילים עבריות (מספרים/אחוזים/לטינית נשארים)."""
    swap = str.maketrans("()[]{}", ")(][}{")
    out = []
    for line in s.split("\n"):
        toks = line.split(" ")
        out.append(" ".join(t[::-1].translate(swap) if _HEB.search(t) else t for t in reversed(toks)))
    return "\n".join(out)


def _ocr_pdf(path: Path, dpi: int = 300, psm: str = "6") -> str:
    import shutil, subprocess
    if not shutil.which("tesseract"):
        return ""
    try:
        import pypdfium2, io
        out = []
        doc = pypdfium2.PdfDocument(str(path))
        for i in range(min(len(doc), 10)):
            buf = io.BytesIO()
            doc[i].render(scale=dpi / 72, no_smoothtext=dpi > 300, no_smoothpath=dpi > 300).to_pil().convert("RGB" if dpi > 300 else "L").save(buf, "PNG", dpi=(dpi, dpi))
            r = subprocess.run(["tesseract", "-", "-", "-l", "heb+eng", "--psm", psm], input=buf.getvalue(), capture_output=True, timeout=120)
            out.append(r.stdout.decode("utf-8", "ignore"))
        return "\n".join(out)
    except Exception as ex:
        print(f"[extract] ocr {path.name}: {ex!r}", file=sys.stderr)
        return ""


def parse_ocr_table_lines(text: str) -> list[dict]:
    """טבלת מדיניות מתוך OCR (PDF סרוק - עובדי המדינה 2026): שורה לכל אפיק - "מניות 94% 94% 6%-/+ 100%-88% ...".
    גבולות = "X%-Y%"; סטייה = "N%-/+" או "+/-N%"; הצפוי = האחוז האחרון לפני הסטייה/הגבולות שנמצא בתוך הגבולות."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    title = next((l for l in lines if "מסלול" in l), "")
    tm = re.search(r"(מסלול[^(\-–]{2,60})", title)
    name = norm_name(tm.group(1)) if tm else "כללי"
    fm = re.search(r"\((\d{2,6})\)", title)
    ym = re.search(r"לשנת\s*(20\d\d)", title) or re.search(r"(20\d\d)", title)
    out = []
    for l in lines:
        l = re.sub(r"[\u200e\u200f\u202a-\u202e]", "", l)
        # גבולות "33%-43%" / "17-25%" (ה-% הראשון נשמט ב-OCR)
        b = re.search(r"(?<![\d.])(\d{1,3}(?:\.\d+)?)\s*%?\s*-\s*(\d{1,3}(?:\.\d+)?)\s*%", l)
        lab = re.match(r"^([א-ת\"'״ ,()\-]{2,60}?)\s*(?=[\d%(]|$)", l)
        lab_txt = lab.group(1) if lab and re.search(r"[א-ת]{2}", lab.group(1)) else ""
        if not lab_txt and re.match(r"^\W*(MSCI|ACWI)", l):
            lab_txt = "מניות"  # שורת מדד הייחוס של המניות - התווית נשברה לשורה הקודמת
        if not lab_txt:
            m_ = re.search(r"(?<=[\d%\s])([א-ת][א-ת\"'״ ,()\-]{1,60})$", l)  # שורה בכיוון הפוך - התווית בסוף
            lab_txt = m_.group(1) if m_ and re.search(r"[א-ת]{2}", m_.group(1)) else ""
        if not b or not re.search(r"[א-ת]{2}", lab_txt):
            continue
        lo, hi = sorted((float(b.group(1)), float(b.group(2))))
        if hi > 150:
            continue
        lb = l[:b.start()] + " " * (b.end() - b.start()) + l[b.end():] if b else l  # הסטייה - לא מתוך הגבולות
        tol = (re.search(r"(?<![\d.])(\d{1,2})\s*%\s*(?:-/\+|\+/-)|(?:\+/-|-/\+)\s*(\d{1,2})\s*%", lb)
               or re.search(r"(?<![\d.])(\d{1,2})\s*%\s*\+|\+\s*(\d{1,2})\s*%", lb))
        cut = min(x.start() for x in (b, tol) if x)
        end = max(x.end() for x in (b, tol) if x)
        before = [float(x) for x in re.findall(r"(\d{1,3}(?:\.\d+)?)\s*%", l[:cut])]
        after = [float(x) for x in re.findall(r"(\d{1,3}(?:\.\d+)?)\s*%", l[end:])]
        if len(after) > len(before):  # שורה הפוכה: גבולות, סטייה, הצפוי, (קודמת), בפועל
            exp = next((x for x in after if lo <= x <= hi), None)
            before = list(reversed(after))
        else:
            exp = next((x for x in reversed(before) if lo <= x <= hi), None)
        asset = lab_txt.strip(" -,(")
        out.append({"fund_id": fm.group(1) if fm else None, "track_no": fm.group(1) if fm else None, "track_code": "ocr",
                    "track_name": name, "group": "ocr", "year": ym.group(1) if ym else None, "asset": asset,
                    "asset_key": asset_key(asset), "current_pct": before[0] if len(before) > 2 else None, "expected_pct": exp,
                    "tolerance": (tol.group(1) or tol.group(2)) + "%" if tol else None, "min_pct": lo, "max_pct": hi,
                    "benchmark": None, "policy_text": l[:200], "sheet": "ocr"})
    return out if len(out) >= 2 else []  # שורה אחת בלבד = OCR חלקי (חסר אפיק המניות?) - לא מפרסמים תמונה חלקית


PROSE_RX = re.compile(r"(?:(עד|לפחות|מקסימום|מינימום)\s*(\d{1,3}(?:\.\d+)?)\s*%|(\d{1,3}(?:\.\d+)?)\s*%\s*(לפחות|לכל היותר))"
                      r"[^.%]{0,60}?(?:יושקעו|יושקע|יוחזקו|יהיו)\s+([^.]{3,120})")


def parse_prose_limits(text: str) -> list[dict]:
    """מדיניות במלל בלי טבלה (דן: "עד 10% מסך הכספים המנוהלים יושקעו במניות ..."): עד -> מקסימום, לפחות -> מינימום."""
    if not re.search(r"מדיניות", text):
        return []
    flat = re.sub(r"\s+", " ", text)
    ym = re.search(r"לשנת\s*(20\d\d)", flat)
    out = []
    for m in PROSE_RX.finditer(flat):
        word = m.group(1) or m.group(4)
        val = float(m.group(2) or m.group(3))
        lab = re.sub(r"[\u200e\u200f\u202a-\u202e]", "", m.group(5))
        lab = re.sub(r"אג[\'י״\"]{1,2}ח", 'אג"ח', lab)
        lab = re.split(r"\s+(?:מדד|בדירוג|באמצעות|בבורסה)\b|,|\s-\s", lab)[0].strip()
        lab = re.sub(r"^(ב|ל)(?=\S)", "", lab)[:60]
        if not re.search(r"אג\"ח|פיקדונ|צמוד", lab) and re.search(r"בורסה|ת[\"״י']?א\s*\d+", m.group(5)):
            lab = "מניות"  # "יושקעו בבורסה לניירות ערך ב- ת"א 125" (דן 2022)
        if not lab or val > 100:
            continue
        mx = word in ("עד", "מקסימום", "לכל היותר")
        out.append({"fund_id": None, "track_no": None, "track_code": "prose", "track_name": "כללי",
                    "year": ym.group(1) if ym else None, "asset": lab, "asset_key": asset_key(lab),
                    "current_pct": None, "expected_pct": None, "tolerance": None,
                    "min_pct": None if mx else val, "max_pct": val if mx else None,
                    "benchmark": None, "policy_text": m.group(0)[:200], "group": "prose", "sheet": "prose"})
    return out


EXPO_RX = re.compile(r"חשופ(?:ים|ות|ה)\s+(?:ל|למ|לנכסי\s+)?(.{2,90}?)(?:,|\s)[^.]{0,160}?(?:בשיעור|בשעור)\s+(?:חשיפה\s+)?שלא\s+יפחת\s+מ[-־\s]*(\d{1,3})\s*%"
                     r"\s*ולא\s+יעלה\s+על\s*(\d{1,3})\s*%")


def parse_exposure_prose(rows, sheet=""):
    """גיליון מסלול מתמחה בלי טבלה (איי.די.איי): "נכסי המסלול יהיו חשופים לנכסי אג"ח ... בשיעור שלא יפחת מ-75% ולא יעלה
    על 120%" -> אפיק עם מינימום/מקסימום; מדד הייחוס מהעמודה "מדד ייחוס". שם המסלול: שורת "מסלול ..." או שם הגיליון."""
    cells = [str(c).strip() for r in rows for c in r if c not in (None, "") and str(c).strip()]
    flat = re.sub(r"\s+", " ", " ".join(cells))
    ms = list(EXPO_RX.finditer(flat))
    if not ms:
        return []
    title = next((c for c in cells[:6] if re.match(r"^מסלול\s", c) and len(c) < 60), "") or sheet
    ym = re.search(r"(?:לשנת|ייחוס)\s*(20\d\d)", flat)
    bench = " ".join(c for c in cells if re.search(r"\d+(\.\d+)?\s*%?\s*-\s*\S|-\s*\d+(\.\d+)?\s*%|%\s*-", c) and len(c) < 60)[:200] or None
    out = []
    for m in ms:
        lab = re.sub(r"^(נכסי|ה)(?=\S)", "", m.group(1)).strip()
        lab = re.split(r"\s+(?:בארץ|בישראל|ובחו|לרבות|בכפוף)\b", lab)[0].strip()[:60]
        if len(lab) < 4 or re.search(r"הבאים|^ם\s", lab):  # "חשופים לנכסים הבאים: ..." -> האפיק לפי שם המסלול
            lab = re.sub(r"^מסלול\s+", "", norm_name(title)).replace("אג ח", 'אג"ח')
        lo, hi = float(m.group(2)), float(m.group(3))
        out.append({"fund_id": None, "track_no": None, "track_code": sheet, "track_name": norm_name(title),
                    "group": sheet, "year": ym.group(1) if ym else None, "asset": lab, "asset_key": asset_key(lab),
                    "current_pct": None, "expected_pct": None, "tolerance": None, "min_pct": lo, "max_pct": hi,
                    "benchmark": bench, "policy_text": m.group(0)[:200], "sheet": sheet})
    return out


_SITE_CFG = {}


def _site_cfg(lid):
    if lid not in _SITE_CFG:
        cf = ROOT / "scripts" / "policy" / "sites" / f"{lid}.json"
        _SITE_CFG[lid] = json.loads(cf.read_text("utf-8")) if lid and cf.exists() else {}
    return _SITE_CFG[lid]


def read_doc(path: Path):
    """-> (text, tables[list[rows]], names[list[str]]) - names = שמות גיליונות (ריק ל-PDF/DOCX)"""
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        texts, tables, names = [], [], []
        with pdfplumber.open(path) as pdf:
            for pn, pg in enumerate(pdf.pages, 1):
                t = pg.extract_text() or ""
                rtl = _is_visual_rtl(t)
                t = _fix_rtl(t) if rtl else t
                texts.append(t)
                # כותרת העמוד: שורה עם מסלול/מדיניות/מספר מ.ה (אנליסט: "אנליסט גמל לבני 50-60 (מ.ה 9731)"), לא הערת שוליים
                lines = [l.strip() for l in t.splitlines() if not re.search(r"ESG|GSE|ראו פרסום", l)]
                title = next((l for l in lines if re.search(r"מסלול|מ\.?ה\.?\W{0,3}\d{3,6}", l)), "") \
                    or next((l for l in lines if "מדיניות" in l), "")  # שורת המסלול עדיפה על כותרת כללית (מחוג)
                for tb in pg.extract_tables():
                    if rtl:  # עברית חזותית: מילים הפוכות וסדר עמודות הפוך
                        tb = [[_fix_rtl(c) if isinstance(c, str) else c for c in reversed(r)] for r in tb]
                    tables.append([[title]] + tb if title else tb)
                    names.append(f"page{pn}")
        text = "\n".join(texts)
        if len(re.sub(r"\s", "", text)) < 40:  # PDF סרוק (דן: מכתב מדיניות מסורק) - OCR עברית אם tesseract מותקן
            text = _ocr_pdf(path) or text
        return text, tables, names
    if ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        tables = [[[str(c) if c != "" else "" for c in sh.row_values(i)] for i in range(sh.nrows)]
                  for sh in wb.sheets()]
        return "\n".join(" ".join(c for c in r if c) for t in tables for r in t), tables, [sh.name for sh in wb.sheets()]
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        tables = [[[str(c) if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
                  for ws in wb.worksheets]
        return "\n".join(" ".join(c for c in r if c) for t in tables for r in t), tables, [ws.title for ws in wb.worksheets]
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
        return "\n".join(p.text for p in d.paragraphs), tables, []
    if ext in (".htm", ".html"):  # מנורה: הצהרות כ-HTML (יצוא Word/Excel, לרוב windows-1255)
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(path.read_bytes(), "html.parser")
        tables = []
        for t in soup.find_all("table"):
            rows = []
            for tr in t.find_all("tr"):
                r = []
                for c in tr.find_all(["td", "th"]):  # colspan -> תאים ריקים, שהעמודות יישארו מיושרות
                    r.append(re.sub(r"\s+", " ", c.get_text(" ", strip=True)))
                    r += [""] * (int(c.get("colspan", 1) or 1) - 1 if str(c.get("colspan", 1)).isdigit() else 0)
                rows.append(r)
            if rows:
                tables.append(rows)
        title = soup.title.get_text(" ", strip=True) if soup.title else ""
        return soup.get_text("\n", strip=True), tables, [title or path.stem] * len(tables)
    return "", [], []  # .doc/.xls ישנים: לא נתמך - מדווח כ-unsupported


# ---------- פרסרים לפי מבנה (מזוהים לפי תוכן הגיליון, לא לפי חברה) ----------
BOUNDS = re.compile(r"(-?\d+(?:\.\d+)?)\s*%?\s*-\s*(-?\d+(?:\.\d+)?)\s*%")

def asset_key(name: str):
    """אפיק סטנדרטי (equity/fx) לפי שם האפיק בחברה; יתר האפיקים נשמרים בשמם המקורי בלבד."""
    n = re.sub(r"[\"'״׳\s]+", "", name or "")
    n = re.sub(r"[\d*()]+$", "", n)  # מספר / כוכבית של הערת שוליים צמודים לשם ("מט"ח6", "חשיפה למט"ח4", "מט"ח(*)")
    if n.startswith(("מניות", "חשיפהלמניות", "סהכמניות")):
        return "equity"
    if n in ("מטח", "חשיפהלמטח", "חשיפהלמטבעחוץ", "חשיפהמטבעית", "חשיפהמטחית"):
        return "fx"
    return None


def _num(v):
    try:
        return float(str(v).replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def parse_statement_blocks(rows, sheet=""):
    """מבנה "הצהרת השקעות": בלוקים של [קידוד מסלול | שם המסלול | כותרת | שורה לאפיק | סוף טבלה].
    הבלוק יכול להתחיל בכל עמודה (בקבצים ישנים מוזז), ולכותרת לא תמיד יש תווית "אפיק" - העמודות לפי טקסט הכותרת.
    -> רשומות long (שורה לכל מסלול x אפיק). [] אם המבנה לא מתאים."""
    grid = [list(r) for r in rows]
    starts = [(ri, ci) for ri, r in enumerate(grid) for ci, c in enumerate(r) if isinstance(c, str) and c.strip() == "קידוד מסלול"]
    if not starts:
        return []
    year = next((m.group(1) for r in grid[:4] for c in r if c and (m := re.search(r"(20\d\d)", str(c)))), None)
    out = []
    for n, (ri, c0) in enumerate(starts):
        end = starts[n + 1][0] if n + 1 < len(starts) else len(grid)
        cell = lambda r, c: grid[r][c] if c < len(grid[r]) else None
        code = str(cell(ri, c0 + 1) or "").strip()
        parts = code.split("-")
        cur = {"track_code": code, "legal_id": parts[0],
               "fund_id": str(int(parts[1])) if len(parts) > 1 and parts[1].isdigit() else None,
               "track_no": parts[2] if len(parts) > 2 else None, "track_name": None, "year": year}
        cols, hdr_row = {}, None
        for r in range(ri + 1, end):
            lab = str(cell(r, c0) or "").strip()
            if lab == "שם המסלול":
                cur["track_name"] = norm_name(str(cell(r, c0 + 1) or ""))
                continue
            if hdr_row is None:
                texts = {c: str(cell(r, c) or "") for c in range(c0, len(grid[r]))}
                if any("גבולות" in t for t in texts.values()):
                    for c, t in texts.items():
                        for key, pat in (("current", r"עדכני|ליום"), ("expected", r"צפוי"), ("tol", r"סטי"),
                                         ("bounds", r"גבולות"), ("bench", r"ייחוס")):
                            if re.search(pat, t) and (key not in cols or key == "current"
                                                      and _hdr_date(t) > _hdr_date(texts[cols[key]])):
                                cols[key] = c
                    hdr_row = r
                continue
            if not lab or lab.startswith(("סוף", "סה\"כ", "*", "תחילת")):
                if lab.startswith("סוף טבלה"):
                    break
                continue
            g = lambda key: cell(r, cols[key]) if key in cols else None
            m = BOUNDS.search(str(g("bounds") or ""))
            cur_pct, exp = _num(g("current")), _num(g("expected"))
            if m is None and exp is None and cur_pct is None:
                continue
            name = lab.rstrip("*").strip()
            out.append({**cur, "asset": name, "asset_key": asset_key(name),
                        "current_pct": _as_pct(cur_pct, g("current")),
                        "expected_pct": _as_pct(exp, g("expected")),
                        "tolerance": g("tol"), "min_pct": float(m.group(1)) if m else None,
                        "max_pct": float(m.group(2)) if m else None,
                        "benchmark": (str(g("bench")).replace("\n", " ") if g("bench") else None), "sheet": sheet})
    return out


MH = re.compile(r"מ[\"״]ה:\s*(\d+)\s*-?\s*(.*)")


def parse_mh_blocks(rows, sheet=""):
    """מבנה 'מ"ה: <קוד קופה> - <שם>': שורות מסלולים (כמה מסלולים יכולים לחלוק טבלה אחת) מעל כותרת 'אפיק השקעה';
    עמודות לפי טקסט הכותרת. בלוקים יכולים לשבת זה לצד זה (פנסיה מקיפה | כללית). legal_id מושלם מהאינדקס."""
    grid = [list(r) for r in rows]
    if not any(isinstance(c, str) and MH.search(c) for r in grid for c in r):
        return []
    year = next((m.group(1) for r in grid[:4] for c in r if c and (m := re.search(r"לשנת\s*(20\d\d)", str(c)))), None)
    out = []
    for ri, row in enumerate(grid):
        anchors = [ci for ci, c in enumerate(row) if isinstance(c, str) and c.strip() == "אפיק השקעה"]
        for n, c in enumerate(anchors):
            end = anchors[n + 1] - 1 if n + 1 < len(anchors) else len(row)  # הבלוק הבא מתחיל בעמודת התווית שלו
            cols = {}
            for ci in range(c, end):
                t = str(row[ci] or "")
                for key, pat in (("current", r"ליום|שיעור החשיפה$"), ("expected", r"צפוי"), ("tol", r"טווח\s*סטי"),
                                 ("bounds", r"גבולות"), ("bench", r"ייחוס")):
                    if key not in cols and re.search(pat, t.strip()):
                        cols[key] = ci
            tracks, r2 = [], ri - 1
            while r2 >= 0 and c < len(grid[r2]) and isinstance(grid[r2][c], str) and (m := MH.search(grid[r2][c])):
                tracks.append((m.group(1), norm_name(m.group(2)))); r2 -= 1
            if not tracks or "expected" not in cols:
                continue
            group = str(grid[r2 + 1][c - 1] or "").replace("\n", " ").strip() if c > 0 else ""
            g = lambda rr, key: grid[rr][cols[key]] if key in cols and cols[key] < len(grid[rr]) else None
            rr = ri + 1
            while rr < len(grid) and c < len(grid[rr]) and grid[rr][c] not in (None, "") and not MH.search(str(grid[rr][c])):
                name = str(grid[rr][c]).strip()
                rr += 1
                if name.startswith("סה"):
                    continue
                cur, exp = _num(g(rr - 1, "current")), _num(g(rr - 1, "expected"))
                m = BOUNDS.search(str(g(rr - 1, "bounds") or ""))
                if cur is None and exp is None and m is None:
                    continue
                for fund_id, tname in tracks:
                    out.append({"fund_id": fund_id, "track_no": fund_id, "track_name": tname, "group": group,
                                "year": year, "asset": name, "asset_key": asset_key(name),
                                "current_pct": _as_pct(cur, g(rr - 1, "current")),
                                "expected_pct": _as_pct(exp, g(rr - 1, "expected")),
                                "tolerance": g(rr - 1, "tol") if _num(g(rr - 1, "tol")) is None and g(rr - 1, "tol") != "ריק במקור" else None,
                                "min_pct": float(m.group(1)) if m else None, "max_pct": float(m.group(2)) if m else None,
                                "benchmark": (str(g(rr - 1, "bench")).strip().replace("\n", " ") if g(rr - 1, "bench") not in (None, "ריק במקור") else None),
                                "sheet": sheet})
                if name.startswith("חשיפה למט"):
                    pass
    return out


PLACEHOLDERS = ("תא ללא תוכן", "הגעת לשדה האחרון", "שורה זו ריקה", "שורה זו אחרונה", "ריק במקור")


def _clean(v):
    s = "" if v is None else str(v).strip()
    return "" if any(s.startswith(p) or s == "תא\xa0ריק\xa0במקור" for p in PLACEHOLDERS) else s


def parse_columns_blocks(rows, sheet=""):
    """מבנה 'בלוק-עמודות לכל מסלול': שורת כותרות עם 'אפיק השקעה' ואז, לכל מסלול, עמודות
    [BM | חשיפה ליום | צפוי | מינימום | מקסימום]; שם המסלול + (קוד קופה) בשורה מעל. טווחים כיחס (0.52) ו-'-' = 0."""
    grid = [[_clean(c) for c in r] for r in rows]
    out = []
    for ri, row in enumerate(grid):
        if not row or row[0] != "אפיק השקעה" or not any(c == "מינימום" for c in row):
            continue
        title_row = grid[ri - 1] if ri > 0 else []
        tol_col = next((j for j, c in enumerate(row) if c.startswith("טווח")), None)
        year = next((m.group(1) for c in (grid[0] if grid else []) for m in [re.fullmatch(r"(20\d\d)", c)] if m), None)
        for j, c in enumerate(row):
            if not c.startswith("חשיפה ליום") or j + 3 >= len(row) or row[j + 1][:4] != "שיעו" or row[j + 2] != "מינימום":
                continue
            cands = [title_row[k] for k in (j, j - 1, j - 2) if 0 <= k < len(title_row) and title_row[k]]
            title = next((t for t in cands if re.search(r"\(\d+\)\s*$", t)), "")
            m = re.search(r"\((\d+)\)\s*$", title)
            bm_col = max((k for k, t in enumerate(row) if t == "BM" and k < j), default=None)
            if not m:
                continue
            fund_id, name = m.group(1), norm_name(title[:m.start()])
            for rr in range(ri + 1, len(grid)):
                r = grid[rr]
                lab = r[0] if r else ""
                if lab.startswith(("שורה זו", "הערה")):
                    break
                if not lab or lab.startswith("סה"):
                    continue
                at = lambda k: _num(r[k]) if k < len(r) and r[k] not in ("", "-") else (0.0 if k < len(r) and r[k] == "-" else None)
                cur, exp, lo, hi = at(j), at(j + 1), at(j + 2), at(j + 3)
                if cur is None and exp is None and lo is None and hi is None:
                    continue
                pct = lambda x: None if x is None else round(x * 100, 2)
                out.append({"fund_id": fund_id, "track_no": fund_id, "track_name": name, "group": sheet.strip(),
                            "year": year, "asset": lab, "asset_key": asset_key(lab), "current_pct": pct(cur),
                            "expected_pct": pct(exp), "tolerance": r[tol_col] if tol_col is not None and tol_col < len(r) else None,
                            "min_pct": pct(lo), "max_pct": pct(hi),
                            "benchmark": (r[bm_col].replace("\n", " ").strip() or None) if bm_col is not None and bm_col < len(r) else None,
                            "sheet": sheet})
    return out


_GEM = dict(zip("אבגדהוזחטיכלמנסעפצקרשת", [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100, 200, 300, 400]))


def heb_year(t):
    """שנה עברית בכותרת -> לועזית ('לשנת תשפ"ו' -> 'לשנת 2026'): קרנות המורים כותבות תשפ"ה | תשפ"ו בשתי עמודות.
    תשפ"ו = 5786 = ספט' 2025 עד אוג' 2026 -> 2026 (השנה שבה מסתיימת, כמו שנת המדיניות שלהן)."""
    def rep(m):
        v = sum(_GEM[ch] for ch in m.group(0) if ch in _GEM)
        return str(5000 + v - 3760) if 700 <= v < 900 else m.group(0)
    return re.sub(r"(?<![א-ת])ת[ש-ת][א-ת]?[\"״][א-ת](?![א-ת])", rep, t) if "ת" in t else t


def parse_titled_tables(rows, sheet=""):
    """מבנה כללי 'שורת כותרת + שורה לאפיק' (הפניקס/רעות/האוניברסיטה/כלל/ארם...): תא 'אפיק (ה)השקעה' בכל עמודה,
    עמודות לפי טקסט הכותרת. גבולות: טקסט '38%-50%', או שני מספרים אחרי 'גבולות' (בכל סדר), או עמודות מינימום/מקסימום.
    שם המסלול: שורה מעל הכותרת; אם גנרי ('מדיניות ... צפויה/שנת') - שם הגיליון. בלי קוד מסלול -> התאמה לפי שם."""
    grid = [[_clean(c) for c in r] for r in rows]
    out = []
    year = next((m.group(1) for r in grid[:6] for c in r for m in [re.search(r"(20\d\d)", c)] if m), None)
    hdr_rx = re.compile(r"^אפיק(\s+(ה)?השקעה|\s*$)")  # "אפיק השקעה" או תא "אפיק" לבד (חח"י השתלמות)
    for ri, row in enumerate(grid):
        lc = next((c for c, t in enumerate(row) if hdr_rx.match(t)), None)
        if lc is None:  # בלי תא "אפיק השקעה": עמודת האפיק = העמודה הריקה משמאל לכותרות (סלייס)
            filled = [c for c, t in enumerate(row) if t]
            nxt = grid[ri + 1] if ri + 1 < len(grid) else []
            if (filled and filled[0] > 0 and any("צפוי" in t for t in row) and any(re.search(r"גבולות|סטי", t) for t in row)
                    and filled[0] - 1 < len(nxt) and re.search(r"[א-ת]", nxt[filled[0] - 1])):
                lc = filled[0] - 1
            else:
                continue
        joined = " ".join(row)
        if "גבולות" not in joined and not ("מינימום" in joined and "מקסימום" in joined):
            continue
        # כותרת שנמשכת בשורות הבאות (PDF אנליסט: "שיעור" / "חשיפה" / "צפוי לשנת" / "2026" בשורות נפרדות):
        # שורות טקסט בלי תווית אפיק ובלי מספרים מצטרפות לתאי הכותרת שמעליהן (לא "מינימום/מקסימום" - נבדק בנפרד)
        row = list(row)
        for nx in grid[ri + 1:ri + 4]:
            if (lc < len(nx) and nx[lc]) or not any(nx) or any(re.search(r"\d%|^\d{1,3}(\.\d+)?$|מינימום|מקסימום", t) for t in nx):
                break
            for c, t in enumerate(nx):
                if t and c < len(row) and row[c]:
                    row[c] = f"{row[c]} {t}"
        raw_row, row = row, [heb_year(t) for t in row]  # השנה העברית רק לבחירת העמודות; שנת המדיניות - כמו קודם (raw_row)
        cols, bounds_cands = {}, []
        pol_years = sorted(int(m.group(1)) for t in row for m in [re.search(r"^מדיניות\s*(20\d\d)", t)] if m)
        for c, t in enumerate(row):
            if c <= lc or not t:
                continue
            py = re.search(r"^מדיניות\s*(20\d\d)", t)
            if "גבולות" in t:
                bounds_cands.append(c)
            if "גבולות" in t and "bounds" not in cols:
                cols["bounds"] = c
            elif "גבולות" in t and re.search(r"20\d\d", t) and re.search(r"20\d\d", row[cols["bounds"]]) \
                    and int(re.search(r"(20\d\d)", t).group(1)) > int(re.search(r"(20\d\d)", row[cols["bounds"]]).group(1)):
                cols["bounds"] = c  # "גבולות ... 2018" | "גבולות ... 2019" (אנליסט): השנה המאוחרת
            elif "מינימום" in t and "min" not in cols:
                cols["min"] = c
            elif "מקסימום" in t and "max" not in cols:
                cols["max"] = c
            elif py and len(pol_years) > 1:  # "מדיניות 2025" | "מדיניות 2026" (הפניקס 2026)
                cols["expected" if int(py.group(1)) == pol_years[-1] else "current"] = c
            elif ("צפוי" in t or "מוצהר" in t or py) and ("expected" not in cols or (
                    "מוצהר" in row[cols["expected"]] + t and re.search(r"(20\d\d)", t) and re.search(r"(20\d\d)", row[cols["expected"]])
                    and int(re.search(r"(20\d\d)", t).group(1)) > int(re.search(r"(20\d\d)", row[cols["expected"]]).group(1)))):
                # "מדיניות מוצהרת 2026" (מינהל) = הצפוי; "מדיניות מוצהרת 2022" | "מדיניות 2023 החל..." (עגור) -> השנה המאוחרת
                cols["expected"] = c
            elif "צפוי" in t and "bounds" not in cols and not re.search(r"גבולות|סטי", t) and (
                    "צפוי" not in row[cols["expected"]]  # הקודם נבחר רק לפי שנה ("מוגדר לשנת 2017") - "צפוי" עדיף (עגור)
                    or (re.search(r"צפוי.*20\d\d", row[cols["expected"]], re.S) and re.search(r"20\d\d", t)
                        and int(re.search(r"(20\d\d)", t).group(1)) > int(re.search(r"(20\d\d)", row[cols["expected"]]).group(1)))):
                cols["expected"] = c  # "צפוי לשנת 2024" | "צפוי לשנת 2025" (עובדי המדינה): השנה המאוחרת = המדיניות
            elif re.search(r"סטי", t) and "tol" not in cols:
                cols["tol"] = c
            elif re.search(r"ייחוס|יחוס", t) and "bench" not in cols:
                cols["bench"] = c
            elif re.search(r"שיעור\s+(ה)?חשיפה", t) and re.search(r"(?<![\d.])(20\d\d)\s*$", t) \
                    and not re.search(r"\d{1,2}[./-]\d{1,2}[./-](20)?\d\d|ליום|לתאריך|נכון ל", t) and (
                    "expected" not in cols or (re.search(r"(20\d\d)\s*$", row[cols["expected"]])
                                               and int(re.search(r"(20\d\d)\s*$", t).group(1))
                                               > int(re.search(r"(20\d\d)\s*$", row[cols["expected"]]).group(1)))):
                cols["expected"] = c  # "שיעור חשיפה 2021" (שנה בלבד, בלי תאריך) = הצפוי; שתי שנים -> המאוחרת (תשפ"ה|תשפ"ו)
            elif re.search(r"שיעור\s+(ה)?חשיפה|ליום|לתאריך|עדכני|נכון ל", t) and (
                    "current" not in cols or _hdr_date(t) > _hdr_date(row[cols["current"]])):
                cols["current"] = c  # שתי עמודות "ליום 31.12.2024 | ליום 31.12.2025" (אל על) - המאוחרת
        if len(bounds_cands) > 1 and "expected" in cols and not any(re.search(r"20\d\d", row[c]) for c in bounds_cands):
            # שתי עמודות "גבולות" בלי שנה (מחוג 2026, קלע 2020: [צפוי 2025, גבולות, ..., צפוי 2026, גבולות]):
            # הגבולות של הצפוי שנבחר = עמודת הגבולות הראשונה אחריו
            after = [c for c in bounds_cands if c > cols["expected"]]
            cols["bounds"] = after[0] if after else cols["bounds"]
        sub = grid[ri + 1] if ri + 1 < len(grid) else []
        if "min" not in cols and "max" not in cols:  # "גבולות" ובשורה שמתחת "מינימום | מקסימום"
            mn = next((c for c, t in enumerate(sub) if t.startswith("מינימום")), None)
            mx = next((c for c, t in enumerate(sub) if t.startswith("מקסימום")), None)
            if mn is not None and mx is not None:
                cols["min"], cols["max"] = mn, mx
        if "expected" not in cols and "bounds" not in cols and "min" not in cols:
            continue
        ey = re.search(r"(20\d\d)", raw_row[cols["expected"]]) if "expected" in cols else None
        tyear = ey.group(1) if ey else year  # "שיעור חשיפה צפוי לשנת 2026" = שנת המדיניות (לא תאריך החשיפה בפועל)
        title, code = "", None
        for r2 in range(ri - 1, max(ri - 6, -1), -1):  # שורות תווית מפורשות: "שם מסלול (מ.ה.)" / "קידוד"
            cells = [c for c in grid[r2] if c]
            if len(cells) >= 2 and re.match(r"^(שם\s+(ה)?(מסלול|קופה)|מסלול\b)", cells[0]):
                title = title or cells[1]
            elif len(cells) >= 2 and cells[0].startswith("קידוד"):
                code = code or cells[1]
            for c2, t2 in enumerate(grid[r2]):  # "קוד קופה" והערך בשורה שמתחת (מגדל)
                if re.match(r"^(קוד|מספר)\s+(קופה|מסלול)", t2) and r2 + 1 < len(grid) and c2 < len(grid[r2 + 1]) \
                        and re.fullmatch(r"\d{2,6}", grid[r2 + 1][c2]):
                    code = code or grid[r2 + 1][c2]
        generic = ""
        for r2 in range(ri - 1, max(ri - 5, -1), -1):
            if title:
                break
            cands = [c for c in grid[r2] if c and not c.startswith(("תחילת", "סוף", "קידוד")) and not re.fullmatch(r"[\d./\-%]+", c)
                     and not re.search(r"ייחוס|יחוס|ESG|GSE|ראו פרסום", c)]
            if cands:
                t0 = max(cands, key=len)
                if re.search(r"^הצהרה|מדיניות.*(צפוי|שנת|לשנת)", t0) and "מסלול" not in t0:
                    generic = generic or t0; continue  # כותרת כללית - שם המסלול בשורה שמעליה (חח"י השתלמות)
                title = t0; break
        title = title or generic
        mt = re.search(r"(מסלול[^(\-–]+)", title or "") if re.search(r"מדיניות.*(צפוי|שנת|לשנת)", title or "") else None
        if mt:  # "מדיניות השקעה צפויה - מסלול כללי (מ"ס מ.ה 382) - לשנת 2021" (PDF עובדי המדינה)
            mf = re.search(r"מ\.?ה\D{0,4}(\d{2,6})|\((\d{2,6})\)", title)
            title = mt.group(1).strip() + (f" {mf.group(1) or mf.group(2)}" if mf else "")
        num_title = re.match(r"^מספר\s+מסלול.*?(\d{3,6})", title or "")
        if num_title:  # "מספר מסלול באוצר 9974": השם בשורה שמעל, המספר = קוד
            above = [c for r2 in range(ri - 1, max(ri - 6, -1), -1) for c in grid[r2]
                     if c and not c.startswith(("תחילת", "סוף", "מספר מסלול")) and not re.search(r"מדיניות.*צפוי", c)]
            title = f"{above[0]} {num_title.group(1)}" if above else title
        if not title or re.search(r"מדיניות.*(צפוי|שנת|לשנת)|הצהרה", title):
            title = sheet.strip() or title
        name = norm_name(title)
        claimed = set(cols.values()) | {lc}

        def g(r, k, off=0):
            if k not in cols:
                return ""
            c = cols[k] + off
            v = r[c] if c < len(r) else ""
            if not v and not off:  # תא ממוזג: הערך בעמודה הסמוכה שאין לה כותרת (PDF אנליסט)
                for c2 in (c + 1, c - 1):
                    if 0 <= c2 < len(r) and c2 not in claimed and r[c2]:
                        return r[c2]
            return v
        n0 = len(out)
        raw = []  # (מיקום ב-out, שורת הנתונים) - להסקת עמודת הצפוי כשאין לה כותרת מזוהה
        for r in grid[ri + 1:]:
            lab = r[lc] if lc < len(r) else ""
            if not lab and lc > 0 and lc - 1 not in claimed and lc - 1 < len(r) and re.search(r"[א-ת]", r[lc - 1] or ""):
                lab = r[lc - 1]  # תווית האפיק בעמודה שמימין לכותרת (PDF אנליסט, תאים ממוזגים)
            if lab.startswith("סוף") or hdr_rx.match(lab) or re.match(r"^(קידוד|שם\s+(ה)?(קופה|מסלול)|מסלולים\s)", lab) \
                    or (len(out) > n0 and any("צפוי" in t for t in r) and any(re.search(r"גבולות|סטי", t) for t in r)):
                break  # תחילת הטבלה הבאה (שורת כותרת אחרי שורות נתונים; לא שורת המשך של הכותרת)
            lab = lab.lstrip("*").strip()
            if not lab or lab.startswith("(") or re.fullmatch(r"[\d./\-%]+", lab) or re.fullmatch(r"סה[\"״]?כ(\s+תיק)?", lab):
                continue
            lab = re.sub(r"^סה[\"״]?כ\s+", "", lab)  # "סה"כ מניות" -> "מניות" (מנורה)
            lo = hi = None
            # "1%" הוא 1, לא שבר - לפי עמודה (ארם: גבולות "40%" אבל הצפוי 0.44)
            bpct = any("%" in g(r, k) for k in ("min", "max", "bounds"))
            if "min" in cols and "max" in cols:
                lo, hi = _num(g(r, "min")), _num(g(r, "max"))
            else:
                bt = g(r, "bounds")
                bt = re.sub(r"^\s*-\s*(\d+(?:\.\d+)?)\s*%\s+(\d+(?:\.\d+)?)\s*%\s*$", r"\1%-\2%", bt)  # "-67% 57%" (PDF רום: המקף זז)
                m = BOUNDS.search(bt)
                if m:  # "51%-39%" (עברית: גבוה-נמוך) -> ממוינים
                    lo, hi = sorted((float(m.group(1)), float(m.group(2))))
                else:
                    nums = [x for x in (_num(g(r, "bounds", k)) for k in range(0, 4)) if x is not None][:2]
                    if len(nums) == 2:
                        lo, hi = min(nums), max(nums)
            if lo is not None and hi is not None and max(abs(lo), abs(hi)) <= 1.5 and not bpct:
                lo, hi = round(lo * 100, 2), round(hi * 100, 2)
            cur, exp = _num(g(r, "current")), _num(g(r, "expected"))
            if cur is None and exp is None and lo is None:
                continue
            pct = lambda x, k: None if x is None else round(x * 100, 2) if abs(x) <= 1.5 and "%" not in g(r, k) else x
            fm = re.search(r"(?:^|\s|\()(\d{3,6})\)?\s*$", name)  # קוד קופה בסוף השם ("כלל פנסיה מניות 9647")
            fid = fm.group(1) if fm and not re.fullmatch(r"(19|20)\d\d", fm.group(1)) else None
            if code and code.isdigit() and len(code) >= 20:  # <ח.פ.><אפסים><קופה><מסלול 5 ספרות> (סלייס)
                fid = str(int(code[-5:]))
            elif code and code.isdigit():
                fid = code
            elif code:  # "קידוד": <ח.פ.>-<קופה>-<מסלול>-<...> (אינפיניטי)
                cp = code.split("-")
                if len(cp) > 2 and cp[2].isdigit():
                    fid = cp[2]
                elif len(cp) > 1 and cp[1].isdigit():
                    fid = str(int(cp[1]))
            out.append({"fund_id": fid, "track_no": fid, "track_code": code or f"{sheet.strip()}|{name}", "track_name": name,
                        "group": sheet.strip(), "year": tyear, "asset": lab, "asset_key": asset_key(lab),
                        "current_pct": pct(cur, "current"), "expected_pct": pct(exp, "expected"), "tolerance": g(r, "tol") or None,
                        "min_pct": lo, "max_pct": hi, "benchmark": (g(r, "bench").replace("\n", " ") or None), "sheet": sheet})
            raw.append((len(out) - 1, r))
        if "expected" not in cols and raw:
            # כותרת הצפוי לא זוהתה (PDF רום: "שיעור חשיפה / צפוי לשנת / 2026" בשלוש שורות ובעמודה אחרת מהערכים):
            # עמודה בלי כותרת שהערכים בה (אחוזים) בתוך הגבולות ברוב השורות = הצפוי
            def inside(c):
                hits = n = 0
                for i, r in raw:
                    lo, hi = out[i]["min_pct"], out[i]["max_pct"]
                    v = _num(r[c]) if c < len(r) else None
                    if lo is None or hi is None or v is None:
                        continue
                    v = v * 100 if abs(v) <= 1.5 and "%" not in r[c] else v
                    n += 1; hits += lo <= v <= hi
                return hits, n
            cands = [(inside(c), c) for c in range(max(len(r) for _, r in raw)) if c not in claimed]
            best = max(cands, default=((0, 0), None))
            (hits, n), bc = best
            if bc is not None and hits >= 2 and hits >= 0.6 * n:
                for i, r in raw:
                    v = _num(r[bc]) if bc < len(r) else None
                    if v is not None and out[i]["expected_pct"] is None:
                        out[i]["expected_pct"] = round(v * 100, 2) if abs(v) <= 1.5 and "%" not in r[bc] else v
    return out


def parse_shape_rows(rows, sheet=""):
    """מוצא אחרון: שורות בצורה [אפיק | % ... | "+/- N%" | גבולות | מדד] כשהכותרת לא מיושרת לעמודות (PDF רום: כותרת
    בשלוש שורות עם פחות תאים מהנתונים). הצפוי = האחוז שלפני הסטייה, בפועל = האחוז שלפניו (אם יש). שם = שורה עם "מסלול" מעל."""
    grid = [[_clean(c) for c in r] for r in rows]
    tol_rx = re.compile(r"^(?:\+/-|-/\+|±)\s*(\d{1,2}(?:\.\d+)?)\s*%$|^(\d{1,2}(?:\.\d+)?)\s*%\s*(?:\+/-|-/\+)$")
    rng_rx = re.compile(r"^-?\s*(\d{1,3}(?:\.\d+)?)\s*%\s*-?\s*(\d{1,3}(?:\.\d+)?)\s*%$")
    pct_rx = re.compile(r"^(\d{1,3}(?:\.\d+)?)\s*%$")
    out, title = [], ""
    for r in grid:
        cells = [c for c in r if c]
        if not cells:
            continue
        if len(cells) == 1 and "מסלול" in cells[0]:
            title = cells[0]; continue
        ti = next((i for i, c in enumerate(cells) if tol_rx.match(c)), None)
        if ti is None or ti + 1 >= len(cells) or not rng_rx.match(cells[ti + 1]) or not re.search(r"[א-ת]{2}", cells[0]):
            continue
        pcts = [float(pct_rx.match(c).group(1)) for c in cells[1:ti] if pct_rx.match(c)]
        if not pcts or not title:
            continue
        m = rng_rx.match(cells[ti + 1]); lo, hi = sorted((float(m.group(1)), float(m.group(2))))
        tm = tol_rx.match(cells[ti])
        ym = re.search(r"(20\d\d)", title)
        nm = re.search(r"(מסלול.+)$", title)
        name = norm_name(nm.group(1)) if nm else norm_name(title)
        out.append({"fund_id": None, "track_no": None, "track_code": f"{sheet.strip()}|{name}", "track_name": name,
                    "group": sheet.strip(), "year": ym.group(1) if ym else None, "asset": cells[0], "asset_key": asset_key(cells[0]),
                    "current_pct": pcts[-2] if len(pcts) > 1 else None, "expected_pct": pcts[-1],
                    "tolerance": (tm.group(1) or tm.group(2)) + "%", "min_pct": lo, "max_pct": hi,
                    "benchmark": " ".join(cells[ti + 2:]).replace("\n", " ") or None, "sheet": sheet})
    return out if len(out) >= 2 else []


def parse_text_tracks(rows, sheet=""):
    """גיליון מסלולים מתמחים מילולי: [שם מסלול (קוד) | מדיניות השקעות (טקסט) | מדד ייחוס]. בלי טווחים מספריים."""
    grid = [[_clean(c) for c in r] for r in rows]
    hdr = next((ri for ri, r in enumerate(grid) if any(c.startswith("מדיניות השקעות") for c in r)
                and any("ייחוס" in c or "בנצ" in c for c in r)), None)
    if hdr is None:
        return []
    h = grid[hdr]
    pc = next(k for k, c in enumerate(h) if c.startswith("מדיניות השקעות"))
    bc = next(k for k, c in enumerate(h) if "ייחוס" in c or "בנצ" in c)
    year = next((m.group(1) for r in grid[:3] for c in r for m in [re.fullmatch(r"(20\d\d)", c)] if m), None)
    out = []
    for r in grid[hdr + 1:]:
        m = re.search(r"^(.*)\((\d+)\)\s*$", r[0]) if r else None
        if not m:
            continue
        text = r[pc] if pc < len(r) else ""
        eq = re.search(r"מניות[^.%]{0,80}?(\d+(?:\.\d+)?)\s*%[^.%]{0,20}?(?:ל|עד|-)\s*-?(\d+(?:\.\d+)?)\s*%", text)
        out.append({"fund_id": m.group(2), "track_no": m.group(2), "track_name": norm_name(m.group(1)), "group": sheet.strip(),
                    "year": year, "asset": "מדיניות מילולית", "asset_key": "equity" if eq else None,
                    "min_pct": float(eq.group(1)) if eq else None, "max_pct": float(eq.group(2)) if eq else None,
                    "current_pct": None, "expected_pct": None, "tolerance": None,
                    "benchmark": (r[bc].replace("\n", " ") if bc < len(r) else None), "policy_text": text[:1500], "sheet": sheet})
    return out


def _hdr_date(t):
    """תאריך בכותרת עמודה ("ליום 31.12.2025", "ליום 30/11/25") כ-YYYY-MM-DD להשוואה; בלי תאריך - ""."""
    m = re.search(r"(\d{1,2})[./](\d{1,2})[./](20\d\d|\d\d)(?!\d)", str(t or ""))
    if not m:
        return ""
    y = m.group(3) if len(m.group(3)) == 4 else "20" + m.group(3)
    return f"{y}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"


def _as_pct(v, raw):
    """שבר (0.48) -> אחוזים; ערך שכבר באחוזים - טקסט עם % ("48.21%", מור) או מעל 1.5 - כמו שהוא (לא 4821)."""
    if v is None:
        return None
    return round(v, 2) if "%" in str(raw) or abs(v) > 1.5 else round(v * 100, 2)


def _iso_date(v):
    """'2026-09-29' / '29/09/2026' / datetime -> 'YYYY-MM-DD', אחרת None."""
    if re.fullmatch(r"4\d{4}(\.0)?", str(v or "").strip()):  # מספר סידורי של אקסל (xls: 46235.0 = 2026-08-01)
        from datetime import date, timedelta
        return (date(1899, 12, 30) + timedelta(days=int(float(v)))).isoformat()
    v = str(v or "")[:10]
    m = re.match(r"(20\d\d)-(\d\d)-(\d\d)", v) or None
    if m:
        return v
    m = re.match(r"(\d{1,2})[./](\d{1,2})[./](20\d\d|\d\d)(?!\d)", v)  # "29/09/2026", "5.11.25"
    return f"{m.group(3) if len(m.group(3)) == 4 else '20' + m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None


_DATE = r"(\d{1,2}[./]\d{1,2}[./](?:20)?\d\d)(?!\d)"
_DOC_DATE_RX = re.compile(
    r"(?:(?:אושר|אושרה|עודכן|עודכנה|מעודכנ?ת?|עדכון)\s*(?:ב?דירקטוריון\s*)?(?:ביום|מיום|בתאריך|מתאריך|החל\s*מ-?)"
    r"|החל\s*מ(?:תאריך|יום)?\s*-?)\s*:?\s*" + _DATE)


def doc_text_date(text, tables):
    """תאריך שכתוב בגוף המסמך על המסמך עצמו: "אושרה בדירקטוריון ביום 26.05.2021" (עובדי המדינה), "עודכן ביום ...",
    "החל מ-...". המאוחר שאינו בעתיד, כ-YYYY-MM-DD. לא "שיעור חשיפה ליום/נכון לתאריך X": בדיקה על 348 מסמכים
    עם תאריך באתר - חציון 143 יום לפני מועד הגרסה (עדכוני אמצע שנה שומרים את עמודת הגרסה השנתית)."""
    hay = [text or ""] + [" ".join(str(c) for c in r if c not in (None, "")) for t in tables for r in t[:40]]
    today = datetime.now().strftime("%Y-%m-%d")
    ds = sorted(d for h in hay for m in _DOC_DATE_RX.finditer(h) if (d := _iso_date(m.group(1))) and d <= today)
    return ds[-1] if ds else None


def parse_change_log(rows, sheet=""):
    """גיליון "מהות שינויים": עמודות לפי שם (מיטב: מספר מסלול | שם מסלול | מהות השינוי | תאריך עדכון;
    אינפיניטי: קופה | שם מסלול | אפיק | מהות השינוי | תאריך עדכון | מספר מגבלה)."""
    if not rows:
        return []
    head = [str(c or "").strip() for c in rows[0]]
    col = lambda *ks: next((i for i, h in enumerate(head) if any(k in h for k in ks)), None)
    c_chg, c_date = col("מהות השינוי"), col("תאריך עדכון")
    if c_chg is None or c_date is None:
        return []
    c_code, c_name, c_asset = col("מספר מסלול", "קידוד"), col("שם מסלול"), col("אפיק")
    get = lambda r, c: r[c] if c is not None and c < len(r) else None
    out = []
    for r in rows[1:]:
        if not r or not any(r):
            continue
        chg = str(get(r, c_chg) or "").replace("\n", " ").strip()
        if c_asset is not None and get(r, c_asset):
            chg = f"{get(r, c_asset)}: {chg}"
        out.append({"track_code": str(get(r, c_code) or get(r, c_name) or ""), "track_name": get(r, c_name),
                    "change": chg, "updated": _iso_date(get(r, c_date)) or str(get(r, c_date) or "")[:10]})
    return out


LONG_FIELDS = ["legal_id", "fund_id", "track_no", "track_code", "track_name", "year", "asset", "asset_key",
               "current_pct", "expected_pct", "tolerance", "min_pct", "max_pct", "benchmark", "policy_text", "group", "doc_first_seen", "url", "doc_file", "sheet"]
CHANGE_FIELDS = ["legal_id", "track_code", "track_name", "change", "updated", "url"]


def dump_layout(names, tables, ent, url):
    """מבנה שאף פרסר לא זיהה: שומרים תחילת כל גיליון כדי שנוסיף פרסר אחרי שנראה אותו."""
    return {"url": url, "legal_id": ent["legal_id"], "file": ent["file"],
            "sheets": [{"name": n, "head": [[str(c)[:60] for c in r if c not in (None, "")] for r in t[:15]]}
                       for n, t in zip(names or [""] * len(tables), tables)]}


FIELDS = ["legal_id", "url", "doc_file", "track_name", "confidence",
          *[f"{k}_{b}" for k in ASSETS for b in ("min", "max")], "benchmark", "snippet"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--all", action="store_true")
    a = ap.parse_args()
    idx_path = OUT / "docs_index.json"
    index = json.loads(idx_path.read_text("utf-8"))
    rows_path = OUT / "tracks_policy.json"
    keep = not a.all and rows_path.exists()
    rows = {r["url"] + "|" + r["track_name"]: r for r in json.loads(rows_path.read_text("utf-8"))} if keep else {}
    long_path, chg_path, unp_path = (OUT / "tracks_policy_long.json", OUT / "policy_changes.json",
                                     OUT / "unparsed_layouts.json")
    _load = lambda p: json.loads(p.read_text("utf-8")) if keep and p.exists() else []
    long_rows, changes, unparsed = _load(long_path), _load(chg_path), _load(unp_path)
    # מסמך שהוצא מהאינדקס (exclude / שיוך שגוי) - גם השורות שלו יוצאות
    long_rows, changes, unparsed = ([r for r in x if r["url"] in index] for x in (long_rows, changes, unparsed))
    rows = {k: r for k, r in rows.items() if r["url"] in index}
    # הגנות זמן (ריצה בענן נקטעת ב-55 דק' ולא נשמר כלום - כל יום מאפס): מסמך שנתקע > POLICY_DOC_TIMEOUT שניות מדולג
    # ומסומן parse_error (לא ינוסה שוב עד שינוי בפרסר); אחרי POLICY_PARSE_BUDGET שניות עוצרים ושומרים - ההרצה הבאה ממשיכה
    import signal, time as _time

    class _DocTimeout(BaseException):  # לא Exception - אחרת "except Exception" בקריאת הקובץ בולע אותו
        pass

    def _alarm(*_):
        raise _DocTimeout()

    doc_timeout = int(os.environ.get("POLICY_DOC_TIMEOUT", "120"))
    stop_at = _time.monotonic() + float(os.environ.get("POLICY_PARSE_BUDGET", "1800"))
    can_alarm = hasattr(signal, "SIGALRM")
    if can_alarm:
        signal.signal(signal.SIGALRM, _alarm)
    done, cur, budget_hit = set(), {}, False
    while True:
        try:
            for url, ent in index.items():
                if url in done:
                    continue
                done.add(url)
                if not a.all and ent.get("parsed_sha") == ent["sha256"] and ent.get("parser_version") == PARSER_VERSION:
                    continue
                if _time.monotonic() > stop_at:
                    budget_hit = True
                    print(f"[extract] parse budget reached - stopping, the next run continues", flush=True)
                    break
                cur.update(url=url, ent=ent)
                if can_alarm:
                    signal.alarm(doc_timeout)
                # החלפת תוצאות קודמות של אותו url (מסמך שהתעדכן)
                long_rows = [r for r in long_rows if r["url"] != url]
                changes = [r for r in changes if r["url"] != url]
                unparsed = [r for r in unparsed if r["url"] != url]
                p = ROOT / ent["file"]
                try:
                    text, tables, names = read_doc(p)
                except Exception as ex:
                    print(f"[extract] {p.name}: {ex!r}", file=sys.stderr); continue
                names = names or [""] * len(tables)
                if _site_cfg(ent["legal_id"]).get("join_split_letter"):
                    # PDF שבו האות האחרונה של מילה נפרדת ("אפיק השקע ה", "מסלול רום הלכ ה" - רום) - רק באתרים שסומנו,
                    # כי בשאר המקומות אות בודדת היא מילה ("קרן ט")
                    fx = lambda v: re.sub(r"(?<=[א-ת]{2}) ([א-ת])(?=[\s\"'״)]|$)", r"\1", v) if isinstance(v, str) else v
                    tables = [[[fx(c) for c in r] for r in t] for t in tables]
                    text = fx(text)
                n_long, leftovers = 0, []
                for nm, t in zip(names, tables):
                    found = parse_statement_blocks(t, nm) or parse_mh_blocks(t, nm) or parse_columns_blocks(t, nm) or parse_titled_tables(t, nm) or parse_text_tracks(t, nm) or parse_shape_rows(t, nm)
                    if found:
                        for r in found:
                            r.setdefault("legal_id", ent["legal_id"])
                            r["track_code"] = r.get("track_code") if r.get("track_code") and r.get("fund_id") is None and "|" in r["track_code"] and r["track_code"].startswith(ent["legal_id"]) else (
                                r.get("track_code") if r.get("fund_id") is not None and r.get("track_code") else
                                f"{ent['legal_id']}|{r['track_code']}" if r.get("track_code") else f"{ent['legal_id']}-{r['fund_id']}")
                            r.update(url=url, doc_file=ent["file"], doc_first_seen=ent.get("first_seen"))
                        long_rows += found; n_long += len(found); continue
                    chg = parse_change_log(t, nm)
                    if chg:
                        for r in chg:
                            r.update(legal_id=ent["legal_id"], url=url)
                        changes += chg
                        # תאריך הגרסה מתוך הקובץ (מיטב: "עדכון מספר 2" בלי תאריך באתר ובשרת; בגיליון "מהות שינויים" - "תאריך עדכון")
                        ds = sorted(d for d in (_iso_date(r.get("updated")) for r in chg)
                                    if d and d <= datetime.now().strftime("%Y-%m-%d"))  # "25.02.2027" (שגיאת הקלדה) - לא
                        if ds:
                            ent["content_date"] = ds[-1]
                        continue
                    ex = parse_exposure_prose(t, nm)  # מסלול מתמחה במלל ("חשופים ל... שלא יפחת מ-75% ולא יעלה על 120%")
                    if ex:
                        for r in ex:
                            r.update(legal_id=ent["legal_id"], track_code=f"{ent['legal_id']}|{nm}|{r['track_name']}", url=url,
                                     doc_file=ent["file"], doc_first_seen=ent.get("first_seen"))
                        long_rows += ex; n_long += len(ex); continue
                    leftovers.append((nm, t))
                # שנה חסרה בגיליון (למשל גיליון מתמחים מילולי) -> שנת המסמך: הרוב בגיליונות האחרים, אחרת מתוך שם הקובץ
                doc_rows = long_rows[len(long_rows) - n_long:] if n_long else []
                years = [r["year"] for r in doc_rows if r.get("year")]
                # שנה בשם הקובץ - בלי הקידומת שלנו (12 תווי sha: "4222020f457f_" הוא לא 2020) ובלי מחרוזות hash
                # (מגדל/Wix: "5adfdb_528d2d37ad314229b89356c20197cc79" - ה-"2019" שבתוכה אינו שנה)
                fname = re.sub(r"[0-9a-f]{16,}", " ", re.sub(r"^[0-9a-f]{12}_", "", Path(ent["file"]).name))
                fy = re.search(r"(20[12]\d)", fname)
                if re.fullmatch(r"(?:[0-9a-f]{12}_)?\d{3,5}", Path(ent["file"]).stem):
                    fy = None  # שם הקובץ הוא מספר המסלול (אלטשולר ".../2017.xlsx" = גמל הלכה), לא שנה
                # הטקסט שהמשתמש רואה באתר (הקישור + השורה/הכרטיס סביבו) גובר על שם הקובץ: "לשנת 2026", או שנה יחידה בטקסט
                # ("מדיניות השקעה 2025"); תווית טכנית ("(html)") - לא טקסט גלוי
                lt = "" if re.match(r"^\((html|xhr|network file|config docs|iframe)\)", ent.get("link_text") or "") else (ent.get("link_text") or "")
                # תאריכים ("12.03.2026", "דצמבר 2024", "(מאי 2018)") הם מועד פרסום/עדכון, לא שנת המדיניות - לא נספרים
                lt_years = set(re.findall(r"(?<!\d)(20[12]\d)(?!\d)", re.sub(
                    r"\d{1,2}[./]\d{1,2}[./](20)?\d\d|(ינואר|פברואר|מרץ|מרס|אפריל|מאי|יוני|יולי|אוגוסט|ספטמבר|אוקטובר|נובמבר|דצמבר)\s*,?\s*(20)?\d\d",
                    " ", lt)))
                fy = re.search(r"לשנת\s*(20[12]\d)", lt) or (re.search(r"(?<!\d)(20[12]\d)(?!\d)", " ".join(lt_years)) if len(lt_years) == 1 else None) or fy
                doc_year = fy.group(1) if fy else (max(set(years), key=years.count) if years else None)
                content_year = max(set(years), key=years.count) if years else None
                if fy and content_year and content_year != fy.group(1):
                    # שם הקובץ והתוכן סותרים (מנורה "...לשנת-2025-...-1343.xlsx" בתיקייה 20260126 עם "צפוי לשנת 2026"):
                    # מועד הפרסום באתר מכריע; בלעדיו - שנה מאוחרת בתוכן גוברת (שנות נתוני עבר בגיליון מוקדמות, לא מאוחרות)
                    from .pubdate import published, fits
                    pub, _ = published(url, ent)
                    if pub and fits(content_year, pub) and not fits(fy.group(1), pub):
                        doc_year = content_year
                    elif not pub and int(content_year) > int(fy.group(1)):
                        doc_year = content_year
                for r in doc_rows:  # שנה בשם הקובץ גוברת (בגוף הגיליון מופיעות לפעמים שנים של נתוני עבר)
                    r["year"] = doc_year or r.get("year")
                # עמודת "חשיפה עדכנית" פגומה במקור (כלל "משתתף חכם": "15" בכל האפיקים או *KEY_ERR) - אותו ערך (לא 0) ב-3+ אפיקים
                # של מסלול = לא נתון; החשיפה העדכנית של המסלול נמחקת (הצפוי והגבולות נשארים)
                by_track = {}
                for r in doc_rows:
                    by_track.setdefault((r.get("track_code"), r.get("sheet"), r.get("track_name")), []).append(r)
                for rs in by_track.values():
                    vals = list({str(r.get("asset") or "").strip(): r.get("current_pct") for r in rs  # אפיק שונה (לא שורה כפולה)
                 if r.get("current_pct") not in (None, 0, 0.0)
                 and not str(r.get("asset") or "").strip().startswith("סה")}.values())  # בלי סה"כ (מניות 100 + סה"כ 100)
                    if len(vals) >= 3 and max(vals.count(v) for v in vals) >= 3 and len(set(vals)) <= max(1, len(vals) // 3):
                        for r in rs:
                            r["current_pct"] = None
                fn_code = re.search(r"-(\d{3,6})\.(xlsx?|pdf)$", Path(ent["file"]).name)  # מנורה: קובץ למסלול, הקוד בשם
                if fn_code and 2016 <= int(fn_code.group(1)) <= datetime.now().year + 1:
                    fn_code = None  # שנת מדיניות (רום "...-2025.pdf"), לא מסלול; מספרי מסלול כמו 2009/2013/2015 (מנורה פנסיה) נשארים
                if fn_code and doc_rows and len({r.get("track_code") for r in doc_rows}) == 1 and not any(r.get("fund_id") for r in doc_rows):
                    for r in doc_rows:
                        r["fund_id"] = r["track_no"] = fn_code.group(1)
                if not n_long and _site_cfg(ent["legal_id"]).get("ocr") and p.suffix.lower() == ".pdf":
                    # PDF סרוק עם שכבת טקסט פגומה (מספרי הטבלה בתמונה) - OCR ופרסור שורות
                    ocr_rows = parse_ocr_table_lines(_ocr_pdf(p))
                    # גליפים מצוירים (אין שכבת טקסט ואין תמונה - עובדי המדינה 7635/15404): גם רזולוציה גבוהה ופריסת עמודות, הטוב מבין השניים
                    alt = parse_ocr_table_lines(_ocr_pdf(p, 450, "4"))
                    ocr_rows = alt if len(alt) > len(ocr_rows) else ocr_rows
                    for r in ocr_rows:
                        r.update(legal_id=ent["legal_id"], track_code=f"{ent['legal_id']}|ocr|{r['track_name']}", url=url,
                                 doc_file=ent["file"], doc_first_seen=ent.get("first_seen"))
                    long_rows += ocr_rows; n_long += len(ocr_rows)
                if not n_long and _site_cfg(ent["legal_id"]).get("prose"):  # מדיניות במלל (מכתב, בלי טבלה) - רק באתרים שסומנו
                    prose = parse_prose_limits(text)
                    for r in prose:
                        r.update(legal_id=ent["legal_id"], track_code=f"{ent['legal_id']}|prose", url=url,
                                 doc_file=ent["file"], doc_first_seen=ent.get("first_seen"))
                    long_rows += prose; n_long += len(prose)
                if fy:  # שורות OCR/מלל נוספו אחרי קביעת שנת המסמך - השנה מהקישור/שם הקובץ (דן: "...לשנת 2022")
                    for r in long_rows[len(long_rows) - n_long:]:
                        r["year"] = r.get("year") or fy.group(1)
                recs =[] if n_long else extract_tracks_from_text(text)
                for nm, t in leftovers:
                    recs += extract_tracks_from_table(t)
                for r in recs:
                    r.update(legal_id=ent["legal_id"], url=url, doc_file=ent["file"])
                    rows[url + "|" + r["track_name"]] = r
                if not n_long and not recs:
                    unparsed.append(dump_layout(names, tables, ent, url))
                if not any(parse_change_log(t) for t in tables):  # בלי יומן שינויים - תאריך כתוב בגוף המסמך
                    cd = doc_text_date(text, tables)
                    if cd:
                        ent["content_date"] = cd
                    else:
                        ent.pop("content_date", None)
                ent["parsed_sha"] = ent["sha256"]
                ent["parser_version"] = PARSER_VERSION
                print(f"[extract] {p.name}: long={n_long} heuristic={len(recs)}", flush=True)
                if can_alarm:
                    signal.alarm(0)
            break
        except _DocTimeout:
            u, e = cur["url"], cur["ent"]
            print(f"[extract] {Path(e['file']).name}: parse timeout ({doc_timeout}s) - skipped", flush=True)
            long_rows = [r for r in long_rows if r["url"] != u]
            changes = [r for r in changes if r["url"] != u]
            rows = {k: r for k, r in rows.items() if r["url"] != u}
            e.update(parse_error=f"timeout {doc_timeout}s", parsed_sha=e["sha256"], parser_version=PARSER_VERSION)
    if can_alarm:
        signal.alarm(0)
    idx_path.write_text(json.dumps(index, ensure_ascii=False, indent=1), "utf-8")
    # track_filter (הגדרות החברה): קובץ אחד מכיל כמה קרנות (הוותיקות: Makefet + Mivtachim באותו גיליון) - נשמרות
    # רק השורות של המסלולים של החברה הזו
    tf = {}
    for lid in {e.get("legal_id") for e in index.values()}:
        cf = ROOT / "scripts" / "policy" / "sites" / f"{lid}.json"
        if lid and cf.exists():
            rx = json.loads(cf.read_text("utf-8")).get("track_filter")
            if rx:
                tf[lid] = re.compile(rx)
    # track_from_sheet (הגדרות האתר): שם המסלול בקובץ כללי ("מסלול רגיל") והמסלול האמיתי בשם הגיליון
    # (קרנות מורים: "מוג מקור אשראי ואגח") -> "<גיליון> - <מסלול>". מוגן מהכפלה (השורות נשמרות בין ריצות)
    for r in long_rows:
        cfg = _site_cfg(r.get("legal_id"))
        sh = (r.get("sheet") or "").strip()
        if cfg.get("track_from_sheet") and sh and not re.match(r"(?i)^(page\s*\d+|גיליון\s*\d*|sheet\s*\d*)$", sh) \
                and not (r.get("track_name") or "").startswith(sh):
            r["track_name"] = f"{sh} - {r.get('track_name') or ''}".strip(" -")
        # track_from_link (regex עם קבוצה): קובץ למסלול, ושם המסלול רק בטקסט הקישור (ילין: "... - ילין לפידות -
        # קופת גמל להשקעה מסלול כללי. עודכן ...") והפרסר מוצא שם כללי (page1)
        lrx = cfg.get("track_from_link")
        if lrx and re.match(r"(?i)^(page\s*\d+|גיליון\s*\d*|sheet\s*\d*|כללי)?$", (r.get("track_name") or "").strip()):
            m = re.search(lrx, (index.get(r.get("url")) or {}).get("link_text") or "")
            if m:
                r["track_name"] = re.sub(r"\s+", " ", m.group(1)).strip()
        # track_name_from_file: [[regex על שם הקובץ, שם מסלול]] - קובץ לכל מסלול, השם רק בשם הקובץ (יחד רופאים: "הורד",
        # PDF סרוק) -> השם מוחלף
        for rx, label in cfg.get("track_name_from_file") or []:
            if re.search(rx, unquote(r.get("url") or "")):
                r["track_name"] = label
                break
        # track_from_file: [[regex על שם הקובץ / url / עמוד המקור / טקסט הקישור, תווית]] - כמה קופות עם אותם שמות מסלולים בקבצים נפרדים
        # (איילון: "איילון מסלול כללי" בגמל להשקעה / השתלמות / גמל לחיסכון) -> "<תווית> - <מסלול>"
        for rx, label in cfg.get("track_from_file") or []:
            ent_ = index.get(r.get("url")) or {}
            where = " ".join(unquote(x or "") for x in (r.get("doc_file"), r.get("url"), ent_.get("source_page"), ent_.get("link_text")))
            if re.search(rx, where) and not (r.get("track_name") or "").startswith(label):
                r["track_name"] = f"{label} - {r.get('track_name') or ''}".strip(" -")
                break
    # track_split (הגדרות האתר): [[regex על שם המסלול, [[שם, מספר], ...]]] - מדיניות אחת לכמה מסלולים (הכשרה: "מסלול אג"ח
    # ממשלות בניהול אלטשולר שחם, מיטב, ילין, מור, אנליסט" = 5 מסלולים, אחד לכל מנהל) -> שורה לכל מסלול עם המספר שלו
    split_rows = []
    for r in long_rows:
        sp = next((sp for rx, sp in (_site_cfg(r.get("legal_id")).get("track_split") or []) if re.search(rx, r.get("track_name") or "")), None)
        if not sp:
            split_rows.append(r); continue
        for nm, n in sp:
            split_rows.append({**r, "track_name": nm, "fund_id": str(n), "track_no": str(n),
                               "track_code": f"{r.get('legal_id')}|split|{n}"})
    long_rows = split_rows
    if tf:
        keep_row = lambda r: r.get("legal_id") not in tf or bool(tf[r["legal_id"]].search(r.get("track_name") or ""))
        long_rows = [r for r in long_rows if keep_row(r)]
        changes = [r for r in changes if keep_row(r)]
        rows = {k: r for k, r in rows.items() if keep_row(r)}
    for r in long_rows:
        sane_bounds(r)
    out = sorted(rows.values(), key=lambda r: (r["legal_id"], r["track_name"]))
    rows_path.write_text(json.dumps(out, ensure_ascii=False, indent=1), "utf-8")
    for path, data in ((long_path, long_rows), (chg_path, changes), (unp_path, unparsed)):
        path.write_text(json.dumps(data, ensure_ascii=False, indent=1), "utf-8")
    for name, data, fields in (("tracks_policy", out, FIELDS), ("tracks_policy_long", long_rows, LONG_FIELDS),
                               ("policy_changes", changes, CHANGE_FIELDS)):
        with open(OUT / f"{name}.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fields, extrasaction="ignore")
            w.writeheader(); w.writerows(data)
    print(f"[extract] long={len(long_rows)} changes={len(changes)} unparsed_docs={len(unparsed)}")


def sane_bounds(r):
    """גבולות הפוכים / משובשים מ-PDF מימין לשמאל (לאומי קמ"פ: "42%-30%" -> min 42 max 30; "26%-16%" -> 2016):
    min>max -> החלפה; גבול משובש (מעל 400) -> צפוי ± סטייה (אם יש), אחרת ריק."""
    lo, hi, exp = r.get("min_pct"), r.get("max_pct"), r.get("expected_pct")
    if lo is None or hi is None:
        return
    if lo > hi:
        lo, hi = hi, lo
    if hi > 400 or lo < -400:  # חשיפה מעל 100% לגיטימית (מניות 94-106); 2016 = שני מספרים שהתחברו
        m = re.search(r"(\d+(?:\.\d+)?)", str(r.get("tolerance") or ""))
        tol = float(m.group(1)) if m else None
        tol = tol * 100 if tol is not None and tol <= 1.5 and "%" not in str(r.get("tolerance")) else tol
        lo, hi = (round(exp - tol, 2), round(exp + tol, 2)) if tol is not None and exp is not None else (None, None)
    r["min_pct"], r["max_pct"] = lo, hi


if __name__ == "__main__":
    main()
